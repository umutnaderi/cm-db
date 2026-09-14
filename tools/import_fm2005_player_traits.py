#!/usr/bin/env python3
"""Validate and import FM2005 preferred moves into Retroball SQLite databases.

The supplied workbook does not contain Retroball source person IDs.  Its binary
record offset is therefore joined to the forensic people_core_v240 export, then
checked against birth date, first-name index and second-name index before any
database row is written.  Date-only joins are deliberately rejected.
"""

from __future__ import annotations

import argparse
import bisect
import csv
import hashlib
import json
import re
import sqlite3
import sys
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable


ROOT = Path(__file__).resolve().parents[1]
DATABASE_SLUG = "fm2005_vanilla_original"
MAPPING_VERSION = "fm2005-ppm-v1"
DEFAULT_CONFIG_DIR = ROOT / "config" / "player_traits" / "fm2005_v1"
DEFAULT_PEOPLE_CORE_ZIP = (
    ROOT / "data" / "profile-ready" / "fm2005_vanilla_d1_profile_ready.zip"
)
DEFINITION_COLUMNS = (
    "trait_key",
    "display_name",
    "source_database_slug",
    "source_trait_id",
    "source_name",
    "source_bit_index",
    "mapping_version",
)
ASSIGNMENT_COLUMNS = (
    "database_slug",
    "source_person_id",
    "source_person_unique_id",
    "date_of_birth",
    "trait_key",
    "source_trait_id",
    "source_trait_name",
    "source_bit_index",
    "source_value",
    "mapping_confidence",
    "source_bitmask_hex",
    "source_record_offset",
    "source_person_record_offset",
    "source_person_record_end_offset",
    "mapping_version",
)


SCHEMA_SQL = """
CREATE TABLE IF NOT EXISTS player_trait_definition (
  trait_key TEXT PRIMARY KEY,
  display_name TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS player_trait_source (
  database_slug TEXT NOT NULL,
  source_person_id TEXT NOT NULL,
  trait_key TEXT NOT NULL,
  source_trait_id INTEGER NOT NULL,
  source_trait_name TEXT NOT NULL,
  source_bit_index INTEGER NOT NULL,
  source_value INTEGER NOT NULL,
  mapping_confidence REAL NOT NULL,
  source_bitmask_hex TEXT NOT NULL,
  source_record_offset INTEGER NOT NULL,
  source_person_unique_id TEXT,
  mapping_version TEXT NOT NULL,
  PRIMARY KEY (database_slug, source_person_id, trait_key),
  FOREIGN KEY (trait_key) REFERENCES player_trait_definition(trait_key)
);

CREATE INDEX IF NOT EXISTS idx_player_trait_source_player
  ON player_trait_source(database_slug, source_person_id);
CREATE INDEX IF NOT EXISTS idx_player_trait_source_trait
  ON player_trait_source(database_slug, trait_key);

CREATE TABLE IF NOT EXISTS player_trait_import (
  database_slug TEXT PRIMARY KEY,
  mapping_version TEXT NOT NULL,
  source_artifact_name TEXT NOT NULL,
  source_artifact_sha256 TEXT NOT NULL,
  source_binary_sha256 TEXT,
  source_extracted_on TEXT,
  definition_count INTEGER NOT NULL,
  player_count INTEGER NOT NULL,
  assignment_count INTEGER NOT NULL
);
"""


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalized_text(value: Any) -> str:
    return str(value or "").strip()


def integer(value: Any, label: str) -> int:
    if isinstance(value, bool) or value is None or value == "":
        raise ValueError(f"{label} is missing")
    number = float(value)
    if not number.is_integer():
        raise ValueError(f"{label} must be an integer: {value!r}")
    return int(number)


def optional_integer(value: Any, label: str) -> int | None:
    return None if value is None or value == "" else integer(value, label)


def iso_date(value: Any, label: str) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = normalized_text(value)
    try:
        return date.fromisoformat(text[:10]).isoformat()
    except ValueError as error:
        raise ValueError(f"{label} is not an ISO date: {value!r}") from error


def trait_key(name: str) -> str:
    key = re.sub(r"[^a-z0-9]+", "_", name.casefold()).strip("_")
    if not key:
        raise ValueError(f"Cannot create a stable trait key from {name!r}")
    return key


def normalized_hex(value: Any) -> str:
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError(f"Bitmask is not integral: {value!r}")
        value = int(value)
    text = normalized_text(value).upper()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    text = text.removeprefix("0X")
    if not re.fullmatch(r"[0-9A-F]+", text):
        raise ValueError(f"Invalid hexadecimal bitmask: {value!r}")
    if len(text) > 12:
        raise ValueError(f"FM2005 PPM bitmask exceeds 48 bits: {text}")
    return text.zfill(12)


def workbook_rows(workbook_path: Path, sheet_name: str) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("Reading the source workbook requires openpyxl") from error

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Workbook is missing the {sheet_name!r} sheet")
        sheet = workbook[sheet_name]
        values = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not values:
        return []
    headers = [normalized_text(value) for value in values[0]]
    if not all(headers):
        raise ValueError(f"{sheet_name} has a blank column heading")
    return [
        dict(zip(headers, row))
        for row in values[1:]
        if any(value is not None and value != "" for value in row)
    ]


def dictionary_rows(workbook_path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("Reading the source workbook requires openpyxl") from error

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["PPM Dictionary"]
        values = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    header_index = next(
        (
            index
            for index, row in enumerate(values)
            if normalized_text(row[0]) == "PPM ID (bit N)"
        ),
        None,
    )
    if header_index is None:
        raise ValueError("PPM Dictionary header was not found")
    definitions = []
    for row in values[header_index + 1 :]:
        if row[0] is None:
            continue
        source_trait_id = integer(row[0], "PPM ID")
        source_name = normalized_text(row[1])
        source_bit_index = integer(row[2], f"PPM {source_trait_id} bit index")
        definitions.append(
            {
                "trait_key": trait_key(source_name),
                "display_name": source_name,
                "source_database_slug": DATABASE_SLUG,
                "source_trait_id": source_trait_id,
                "source_name": source_name,
                "source_bit_index": source_bit_index,
                "mapping_version": MAPPING_VERSION,
            }
        )
    expected_ids = list(range(1, 49))
    if [row["source_trait_id"] for row in definitions] != expected_ids:
        raise ValueError("PPM Dictionary must contain source trait IDs 1 through 48 in order")
    if any(row["source_bit_index"] != row["source_trait_id"] - 1 for row in definitions):
        raise ValueError("PPM Dictionary bit indexes do not match the source IDs")
    if len({row["trait_key"] for row in definitions}) != len(definitions):
        raise ValueError("Normalized trait keys are not unique")
    return definitions


def zip_member(archive: zipfile.ZipFile, suffix: str) -> str:
    matches = [name for name in archive.namelist() if name.endswith(suffix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one {suffix} member, found {len(matches)}")
    return matches[0]


def load_people_core(zip_path: Path) -> tuple[list[dict[str, Any]], str | None]:
    with zipfile.ZipFile(zip_path) as archive:
        core_name = zip_member(archive, "/people_core_v240.csv")
        with archive.open(core_name) as raw:
            rows = list(csv.DictReader(line.decode("utf-8-sig") for line in raw))
        overview_name = zip_member(archive, "/source_file_overview.csv")
        with archive.open(overview_name) as raw:
            overview = list(csv.DictReader(line.decode("utf-8-sig") for line in raw))
    people_binary = next(
        (row for row in overview if row.get("filename") == "people_db.dat"),
        None,
    )
    prepared = []
    for row in rows:
        prepared.append(
            {
                **row,
                "file_offset": integer(row["file_offset"], "people_core file_offset"),
                "record_end_offset": integer(
                    row["record_end_offset"], "people_core record_end_offset"
                ),
                "source_person_id": normalized_text(row["source_person_id"]),
                "person_unique_id": normalized_text(row["person_unique_id"]),
                "first_name_id": integer(row["first_name_id"], "people_core first_name_id"),
                "second_name_id": integer(row["second_name_id"], "people_core second_name_id"),
                "dob_days": integer(row["dob_days"], "people_core dob_days"),
                "dob_year": integer(row["dob_year"], "people_core dob_year"),
            }
        )
    prepared.sort(key=lambda row: row["file_offset"])
    if any(
        current["record_end_offset"] > following["file_offset"]
        for current, following in zip(prepared, prepared[1:])
    ):
        raise ValueError("people_core contains overlapping person records")
    return prepared, people_binary.get("sha256") if people_binary else None


def source_birth_date(person: dict[str, Any]) -> str:
    return (
        date(person["dob_year"], 1, 1) + timedelta(days=person["dob_days"])
    ).isoformat()


def split_trait_ids(value: Any) -> list[int]:
    return [integer(part, "Trait ID") for part in normalized_text(value).split(",") if part]


def build_config(workbook_path: Path, people_core_zip: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    definitions = dictionary_rows(workbook_path)
    definition_by_id = {row["source_trait_id"]: row for row in definitions}
    players = workbook_rows(workbook_path, "Players")
    traits_long = workbook_rows(workbook_path, "Traits Long")
    people, source_binary_sha256 = load_people_core(people_core_zip)
    starts = [row["file_offset"] for row in people]

    assignments: list[dict[str, Any]] = []
    mapped_players: dict[int, dict[str, Any]] = {}
    for row_number, row in enumerate(players, start=2):
        record_offset = integer(row.get("Record offset"), f"Players row {row_number} Record offset")
        index = bisect.bisect_right(starts, record_offset) - 1
        if index < 0:
            raise ValueError(f"Players row {row_number} precedes the first person record")
        person = people[index]
        if not person["file_offset"] <= record_offset < person["record_end_offset"]:
            raise ValueError(
                f"Players row {row_number} offset {record_offset} is outside a person record"
            )

        dob = iso_date(row.get("DOB"), f"Players row {row_number} DOB")
        if dob != source_birth_date(person):
            raise ValueError(
                f"Players row {row_number} DOB {dob} does not match person "
                f"{person['source_person_id']} ({source_birth_date(person)})"
            )
        index_a = optional_integer(row.get("Index A"), f"Players row {row_number} Index A")
        index_b = optional_integer(row.get("Index B"), f"Players row {row_number} Index B")
        if index_a is not None and index_a != person["first_name_id"]:
            raise ValueError(
                f"Players row {row_number} Index A {index_a} does not match first-name "
                f"index {person['first_name_id']}"
            )
        if index_b is not None and index_b != person["second_name_id"]:
            raise ValueError(
                f"Players row {row_number} Index B {index_b} does not match second-name "
                f"index {person['second_name_id']}"
            )
        if record_offset in mapped_players:
            raise ValueError(f"Workbook repeats source record offset {record_offset}")
        mapped_players[record_offset] = person

        trait_ids = split_trait_ids(row.get("Trait IDs"))
        if integer(row.get("Trait count"), f"Players row {row_number} Trait count") != len(trait_ids):
            raise ValueError(f"Players row {row_number} trait count is inconsistent")
        if len(set(trait_ids)) != len(trait_ids):
            raise ValueError(f"Players row {row_number} repeats a trait ID")
        mask_hex = normalized_hex(row.get("Bitmask hex"))
        expected_mask = sum(1 << (trait_id - 1) for trait_id in trait_ids)
        if int(mask_hex, 16) != expected_mask:
            raise ValueError(f"Players row {row_number} bitmask does not match its trait IDs")
        names = [name.strip() for name in normalized_text(row.get("Preferred Moves")).split(";")]
        expected_names = [definition_by_id[trait_id]["source_name"] for trait_id in trait_ids]
        if names != expected_names:
            raise ValueError(f"Players row {row_number} trait names do not match the dictionary")

        for trait_id in trait_ids:
            definition = definition_by_id.get(trait_id)
            if not definition:
                raise ValueError(f"Players row {row_number} has unknown trait ID {trait_id}")
            assignments.append(
                {
                    "database_slug": DATABASE_SLUG,
                    "source_person_id": person["source_person_id"],
                    "source_person_unique_id": person["person_unique_id"],
                    "date_of_birth": dob,
                    "trait_key": definition["trait_key"],
                    "source_trait_id": trait_id,
                    "source_trait_name": definition["source_name"],
                    "source_bit_index": definition["source_bit_index"],
                    "source_value": 1,
                    "mapping_confidence": "1.0",
                    "source_bitmask_hex": mask_hex,
                    "source_record_offset": record_offset,
                    "source_person_record_offset": person["file_offset"],
                    "source_person_record_end_offset": person["record_end_offset"],
                    "mapping_version": MAPPING_VERSION,
                }
            )

    if len({row["source_person_id"] for row in assignments}) != len(players):
        raise ValueError("Two workbook player records mapped to the same source person")

    expected_long = sorted(
        (
            row["source_record_offset"],
            row["source_trait_id"],
            row["source_trait_name"],
            row["date_of_birth"],
        )
        for row in assignments
    )
    actual_long = sorted(
        (
            integer(row.get("Record offset"), "Traits Long Record offset"),
            integer(row.get("PPM ID"), "Traits Long PPM ID"),
            normalized_text(row.get("Preferred Move")),
            iso_date(row.get("DOB"), "Traits Long DOB"),
        )
        for row in traits_long
    )
    if actual_long != expected_long:
        raise ValueError("Traits Long does not exactly match the Players sheet expansion")

    roberto = [
        row
        for row in assignments
        if row["source_person_id"] == "4357"
    ]
    if (
        {row["source_trait_id"] for row in roberto} != {9, 10, 15, 37, 46}
        or {row["source_person_unique_id"] for row in roberto} != {"11944"}
    ):
        raise ValueError("Verified Roberto Carlos sentinel did not map exactly")

    manifest = {
        "schema": "retroball.player-traits-source",
        "version": 1,
        "database_slug": DATABASE_SLUG,
        "mapping_version": MAPPING_VERSION,
        "source_artifact_name": workbook_path.name,
        "source_artifact_sha256": sha256_file(workbook_path),
        "source_people_core_zip_sha256": sha256_file(people_core_zip),
        "source_binary_sha256": source_binary_sha256,
        "source_extracted_on": "2026-09-14",
        "definition_count": len(definitions),
        "player_count": len(players),
        "assignment_count": len(assignments),
        "identity_join": {
            "method": "binary record-offset containment",
            "required_checks": [
                "record boundary",
                "date of birth",
                "first-name index when present",
                "second-name index when present",
            ],
            "unmatched_players": 0,
            "ambiguous_players": 0,
            "mapping_confidence": 1.0,
        },
        "verified_sentinel": {
            "name": "Roberto Carlos",
            "source_person_id": "4357",
            "source_person_unique_id": "11944",
            "source_trait_ids": [9, 10, 15, 37, 46],
        },
    }
    return definitions, assignments, manifest


def write_csv(path: Path, columns: Iterable[str], rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=list(columns), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def write_config(
    config_dir: Path,
    definitions: list[dict[str, Any]],
    assignments: list[dict[str, Any]],
    manifest: dict[str, Any],
) -> dict[str, Any]:
    definitions_path = config_dir / "trait_definitions.csv"
    assignments_path = config_dir / "player_traits.csv"
    manifest_path = config_dir / "manifest.json"
    write_csv(definitions_path, DEFINITION_COLUMNS, definitions)
    write_csv(assignments_path, ASSIGNMENT_COLUMNS, assignments)
    manifest = {
        **manifest,
        "trait_definitions_sha256": sha256_file(definitions_path),
        "player_traits_sha256": sha256_file(assignments_path),
    }
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return manifest


def load_config(config_dir: Path) -> tuple[list[dict[str, str]], list[dict[str, str]], dict[str, Any]]:
    definitions_path = config_dir / "trait_definitions.csv"
    assignments_path = config_dir / "player_traits.csv"
    manifest_path = config_dir / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if sha256_file(definitions_path) != manifest.get("trait_definitions_sha256"):
        raise ValueError("trait_definitions.csv does not match its manifest hash")
    if sha256_file(assignments_path) != manifest.get("player_traits_sha256"):
        raise ValueError("player_traits.csv does not match its manifest hash")
    with definitions_path.open(encoding="utf-8", newline="") as source:
        definitions = list(csv.DictReader(source))
    with assignments_path.open(encoding="utf-8", newline="") as source:
        assignments = list(csv.DictReader(source))
    if len(definitions) != int(manifest["definition_count"]):
        raise ValueError("Trait definition count does not match the manifest")
    if len(assignments) != int(manifest["assignment_count"]):
        raise ValueError("Trait assignment count does not match the manifest")
    if len({row["source_person_id"] for row in assignments}) != int(manifest["player_count"]):
        raise ValueError("Trait player count does not match the manifest")
    return definitions, assignments, manifest


def validate_database_players(
    connection: sqlite3.Connection,
    assignments: list[dict[str, str]],
) -> None:
    assignment_dobs: dict[str, str] = {}
    for row in assignments:
        person_id = row["source_person_id"]
        previous = assignment_dobs.setdefault(person_id, row["date_of_birth"])
        if previous != row["date_of_birth"]:
            raise ValueError(f"Config has conflicting birth dates for person {person_id}")
    profiles: dict[str, str] = {}
    person_ids = sorted(assignment_dobs)
    for start in range(0, len(person_ids), 500):
        batch = person_ids[start : start + 500]
        placeholders = ",".join("?" for _ in batch)
        profiles.update(
            {
                str(row[0]): normalized_text(row[1])
                for row in connection.execute(
                    f"""
                    SELECT source_person_id, date_of_birth
                    FROM player_profile
                    WHERE database_slug = ?
                      AND source_person_id IN ({placeholders})
                    """,
                    (DATABASE_SLUG, *batch),
                )
            }
        )
    missing = sorted(set(assignment_dobs) - set(profiles))
    if missing:
        raise ValueError(f"Database is missing {len(missing)} mapped players; first: {missing[0]}")
    mismatched = [
        person_id
        for person_id, dob in assignment_dobs.items()
        if profiles[person_id] != dob
    ]
    if mismatched:
        person_id = mismatched[0]
        raise ValueError(
            f"Database birth date mismatch for person {person_id}: "
            f"database={profiles[person_id]!r}, config={assignment_dobs[person_id]!r}"
        )


def import_database(
    database_path: Path,
    definitions: list[dict[str, str]],
    assignments: list[dict[str, str]],
    manifest: dict[str, Any],
    apply: bool,
) -> dict[str, Any]:
    if not database_path.is_file():
        raise FileNotFoundError(database_path)
    connection = sqlite3.connect(database_path)
    try:
        validate_database_players(connection, assignments)
        if not apply:
            return {
                "database": str(database_path.resolve()),
                "mode": "dry-run",
                "validated_players": manifest["player_count"],
                "assignments": manifest["assignment_count"],
            }

        connection.executescript(SCHEMA_SQL)
        connection.execute("BEGIN IMMEDIATE")
        connection.executemany(
            """
            INSERT INTO player_trait_definition (trait_key, display_name)
            VALUES (?, ?)
            ON CONFLICT(trait_key) DO UPDATE SET display_name = excluded.display_name
            """,
            [(row["trait_key"], row["display_name"]) for row in definitions],
        )
        connection.execute(
            "DELETE FROM player_trait_source WHERE database_slug = ?",
            (DATABASE_SLUG,),
        )
        connection.executemany(
            """
            INSERT INTO player_trait_source (
              database_slug, source_person_id, trait_key, source_trait_id,
              source_trait_name, source_bit_index, source_value,
              mapping_confidence, source_bitmask_hex, source_record_offset,
              source_person_unique_id, mapping_version
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    row["database_slug"],
                    row["source_person_id"],
                    row["trait_key"],
                    int(row["source_trait_id"]),
                    row["source_trait_name"],
                    int(row["source_bit_index"]),
                    int(row["source_value"]),
                    float(row["mapping_confidence"]),
                    row["source_bitmask_hex"],
                    int(row["source_record_offset"]),
                    row["source_person_unique_id"],
                    row["mapping_version"],
                )
                for row in assignments
            ],
        )
        connection.execute(
            """
            INSERT INTO player_trait_import (
              database_slug, mapping_version, source_artifact_name,
              source_artifact_sha256, source_binary_sha256, source_extracted_on,
              definition_count, player_count, assignment_count
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(database_slug) DO UPDATE SET
              mapping_version = excluded.mapping_version,
              source_artifact_name = excluded.source_artifact_name,
              source_artifact_sha256 = excluded.source_artifact_sha256,
              source_binary_sha256 = excluded.source_binary_sha256,
              source_extracted_on = excluded.source_extracted_on,
              definition_count = excluded.definition_count,
              player_count = excluded.player_count,
              assignment_count = excluded.assignment_count
            """,
            (
                DATABASE_SLUG,
                manifest["mapping_version"],
                manifest["source_artifact_name"],
                manifest["source_artifact_sha256"],
                manifest.get("source_binary_sha256"),
                manifest.get("source_extracted_on"),
                manifest["definition_count"],
                manifest["player_count"],
                manifest["assignment_count"],
            ),
        )

        counts = connection.execute(
            """
            SELECT count(*), count(DISTINCT source_person_id)
            FROM player_trait_source
            WHERE database_slug = ?
            """,
            (DATABASE_SLUG,),
        ).fetchone()
        if counts != (manifest["assignment_count"], manifest["player_count"]):
            raise ValueError(f"Imported trait counts are wrong: {counts!r}")
        roberto = [
            row[0]
            for row in connection.execute(
                """
                SELECT source_trait_id
                FROM player_trait_source
                WHERE database_slug = ? AND source_person_id = '4357'
                ORDER BY source_trait_id
                """,
                (DATABASE_SLUG,),
            )
        ]
        if roberto != [9, 10, 15, 37, 46]:
            raise ValueError(f"Roberto Carlos sentinel failed after import: {roberto!r}")
        connection.commit()
        connection.execute("PRAGMA wal_checkpoint(TRUNCATE)")
        return {
            "database": str(database_path.resolve()),
            "mode": "applied",
            "definitions": manifest["definition_count"],
            "players": counts[1],
            "assignments": counts[0],
            "integrity": connection.execute(
                "PRAGMA integrity_check('player_trait_source')"
            ).fetchone()[0],
        }
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, help="Decoded FM2005 traits workbook")
    parser.add_argument(
        "--people-core-zip",
        type=Path,
        default=DEFAULT_PEOPLE_CORE_ZIP,
        help="Profile-ready archive containing people_core_v240.csv",
    )
    parser.add_argument(
        "--config-dir",
        type=Path,
        default=DEFAULT_CONFIG_DIR,
        help="Versioned normalized trait source directory",
    )
    parser.add_argument(
        "--write-config",
        action="store_true",
        help="Validate the workbook and replace the normalized config files",
    )
    parser.add_argument(
        "--database",
        type=Path,
        action="append",
        default=[],
        help="SQLite database to validate or update; may be repeated",
    )
    parser.add_argument(
        "--apply",
        action="store_true",
        help="Apply the import transaction; otherwise database checks are read-only",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        if args.write_config:
            if not args.workbook or not args.workbook.is_file():
                raise ValueError("--write-config requires an existing --workbook")
            if not args.people_core_zip.is_file():
                raise FileNotFoundError(args.people_core_zip)
            definitions, assignments, manifest = build_config(
                args.workbook.resolve(), args.people_core_zip.resolve()
            )
            manifest = write_config(
                args.config_dir.resolve(), definitions, assignments, manifest
            )
        else:
            definitions, assignments, manifest = load_config(args.config_dir.resolve())

        results = [
            import_database(
                database.resolve(), definitions, assignments, manifest, args.apply
            )
            for database in args.database
        ]
        print(json.dumps({"manifest": manifest, "databases": results}, indent=2, sort_keys=True))
        return 0
    except Exception as error:
        print(f"FM2005 trait import failed: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
